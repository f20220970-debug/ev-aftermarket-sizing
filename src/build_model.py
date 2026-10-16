"""Builds model/ev2w_service_market.xlsx: a formula-driven market-sizing model for India's
electric two-wheeler (E2W) servicing and battery-replacement aftermarket, FY22-FY31.

Every number on the Inputs sheet is either sourced (blue, with the source next to it) or a
labelled assumption (blue on yellow). Every other number is a formula, so changing an input
re-sizes the whole market. Re-run this script after editing inputs here.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

OUT = Path(__file__).resolve().parents[1] / "model" / "ev2w_service_market.xlsx"
YEARS = list(range(2022, 2032))
FIRST_COL = 3                                        # column C = FY22
BLUE, BLACK, GREEN = Font(name="Arial", color="0000FF"), Font(name="Arial"), Font(name="Arial", color="008000")
BOLD, TITLE = Font(name="Arial", bold=True), Font(name="Arial", bold=True, size=14)
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1F3A5F")
WHITE_B = Font(name="Arial", bold=True, color="FFFFFF")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
CR = '#,##0;(#,##0);"-"'
UNITS = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'

SOURCES = {
    "argus": "Argus Media, 'India's EV sales hit record high in FY2022-23' (Vahan data): E2W 252,555 units in FY2021-22.",
    "pib": "PIB annexure, state/month-wise E-2W registrations FY2022-23 (Vahan 4): grand total 7,28,271.",
    "mhi": "Ministry of Heavy Industries data via Times Drive: E2W 11,49,334 units in FY25 vs 9,48,561 in FY24.",
    "iesa": "IESA 'India's Mobility Report FY2025-26' via EV Infrastructure News: 2.55 million EV sales in FY2025-26.",
    "jmk": "JMK Research via EV Infrastructure News: E2W nearly 58% of total EV sales in FY2025-26.",
    "evr": "EVreporter (Vahan, 2 Jan 2026): E2W sales 12,79,951 in CY2025, +11.4% YoY; 6.3% of all 2W sales.",
    "meraev": "MeraEV (2026): typical annual e-scooter maintenance Rs 2,000-5,000, about half a 110cc petrol scooter (Rs 3,500-6,000).",
    "battery": "e-vehicleinfo / BikeLeague: e-scooter battery replacement roughly Rs 55,000-60,000.",
    "iesa25": "IESA via IANS (Jan 2026): Karnataka about 9% of national EV sales in CY2025 (all EV types).",
}


def hdr(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, start + i, t)
        c.font, c.fill, c.alignment = WHITE_B, HEAD, Alignment(horizontal="center", wrap_text=True)


def inputs(wb):
    ws = wb.create_sheet("Inputs")
    ws["A1"], ws["A1"].font = "Inputs: sourced data (blue) and assumptions (blue on yellow)", TITLE
    hdr(ws, 3, ["Historical E2W sales", "Value", "Unit", "Source"])
    rows = [("FY22 E2W sales", 252555, "units", SOURCES["argus"]),
            ("FY23 E2W sales", 728271, "units", SOURCES["pib"]),
            ("FY24 E2W sales", 948561, "units", SOURCES["mhi"]),
            ("FY25 E2W sales", 1149334, "units", SOURCES["mhi"]),
            ("FY26 total EV sales", 2550000, "units", SOURCES["iesa"]),
            ("FY26 E2W share of EV sales", 0.58, "%", SOURCES["jmk"])]
    for i, (a, v, u, s) in enumerate(rows, start=4):
        ws.cell(i, 1, a); c = ws.cell(i, 2, v); c.font = BLUE
        c.number_format = PCT if u == "%" else UNITS
        ws.cell(i, 3, u); ws.cell(i, 4, s)
    ws["A10"], ws["B10"], ws["C10"] = "FY26 E2W sales (estimate)", "=B8*B9", "units"
    ws["B10"].number_format = UNITS
    ws["D10"] = "Derived: total FY26 EV sales x E2W share (two different sources, so an estimate)."

    hdr(ws, 12, ["Scenario assumptions", "Bear", "Base", "Bull", "Basis"])
    scen = [("Annual E2W sales growth, FY27-FY31", (0.10, 0.20, 0.30), PCT,
             "Assumption. Context: CY25 growth +11.4% (" + SOURCES["evr"][:9] + "), FY25 +21% (MHI)."),
            ("Service + wear-parts spend per paying vehicle (Rs/yr)", (2000, 3000, 4000), UNITS,
             "Assumption within the Rs 2,000-5,000 range reported by MeraEV."),
            ("Share of vehicles replacing battery at replacement age", (0.40, 0.50, 0.60), PCT,
             "Assumption: many packs are covered by warranty or last longer; others are scrapped instead.")]
    for i, (a, vals, fmt, note) in enumerate(scen, start=13):
        ws.cell(i, 1, a)
        for j, v in enumerate(vals, start=2):
            c = ws.cell(i, j, v); c.font, c.fill, c.number_format = BLUE, YELLOW, fmt
        ws.cell(i, 5, note)

    hdr(ws, 17, ["Common assumptions", "Value", "Unit", "Basis"])
    common = [("Fleet scrappage per year", 0.03, "%", "Assumption: accidents, theft, early failures."),
              ("Battery pack price in FY26", 55000, "Rs", SOURCES["battery"]),
              ("Battery price decline per year", 0.07, "%", "Assumption: falling cell costs."),
              ("Battery replacement age", 6, "years", "Assumption: typical warranty 3-5 years plus wear."),
              ("Independent-garage share of spend", 0.25, "%", "Assumption: rest goes to OEM service networks."),
              ("Karnataka share of national market", 0.09, "%", SOURCES["iesa25"] + " Used as a proxy.")]
    for i, (a, v, u, s) in enumerate(common, start=18):
        ws.cell(i, 1, a); c = ws.cell(i, 2, v); c.font = BLUE
        c.fill = YELLOW if "Assumption" in s else PatternFill()
        c.number_format = PCT if u == "%" else UNITS
        ws.cell(i, 3, u); ws.cell(i, 4, s)
    ws["A25"] = "Pre-FY22 E2W fleet is excluded (small: high-speed E2W sales before FY22 were a fraction of FY22's)."
    for col, w in zip("ABCDE", [52, 14, 14, 14, 110]):
        ws.column_dimensions[col].width = w
    return ws


SCEN_COL = {"Bear": "B", "Base": "C", "Bull": "D"}
BLOCK_START = {"Bear": 4, "Base": 18, "Bull": 32}
LINES = ["E2W sales (units)", "Fleet in use, year end (units)", "Paying fleet (units)",
         "Service + wear-parts market (Rs Cr)", "Battery replacements (units)", "Battery pack price (Rs)",
         "Battery replacement market (Rs Cr)", "Total aftermarket (Rs Cr)",
         "Independent-garage addressable (Rs Cr)", "Karnataka market (Rs Cr)", "Battery share of total"]


def model(wb):
    ws = wb.create_sheet("Model")
    ws["A1"], ws["A1"].font = "Model: fleet and aftermarket by scenario (all formulas)", TITLE
    ws.cell(2, 1, "Fiscal year").font = BOLD
    for k, y in enumerate(YEARS):
        c = ws.cell(2, FIRST_COL + k, y); c.font = WHITE_B; c.fill = HEAD; c.alignment = Alignment(horizontal="center")
    hist = {2022: "Inputs!$B$4", 2023: "Inputs!$B$5", 2024: "Inputs!$B$6", 2025: "Inputs!$B$7", 2026: "Inputs!$B$10"}
    for name, r0 in BLOCK_START.items():
        s = SCEN_COL[name]
        ws.cell(r0, 1, f"{name} scenario").font = BOLD
        g, spend, share = f"Inputs!${s}$13", f"Inputs!${s}$14", f"Inputs!${s}$15"
        for i, label in enumerate(LINES, start=1):
            ws.cell(r0 + i, 1, label)
        for k, y in enumerate(YEARS):
            col = L(FIRST_COL + k); prev = L(FIRST_COL + k - 1)
            r = {n: r0 + i for i, n in enumerate(["sales", "fleet", "pay", "svc", "bu", "bp", "bm", "tot", "ind", "ka", "bsh"], start=1)}
            sales = f"={hist[y]}" if y in hist else f"={prev}{r['sales']}*(1+{g})"
            fleet = f"={col}{r['sales']}" if k == 0 else f"={prev}{r['fleet']}*(1-Inputs!$B$18)+{col}{r['sales']}"
            cells = {
                "sales": (sales, UNITS, GREEN if y in hist else BLACK),
                "fleet": (fleet, UNITS, BLACK),
                "pay": (f"={col}{r['fleet']}-{col}{r['sales']}", UNITS, BLACK),
                "svc": (f"={col}{r['pay']}*{spend}/10^7", CR, BLACK),
                "bu": (f"=IF({col}$2-Inputs!$B$21<{YEARS[0]},0,OFFSET({col}{r['sales']},0,-Inputs!$B$21)*{share}*(1-Inputs!$B$18)^Inputs!$B$21)", UNITS, BLACK),
                "bp": (f"=Inputs!$B$19*(1-Inputs!$B$20)^({col}$2-2026)", UNITS, BLACK),
                "bm": (f"={col}{r['bu']}*{col}{r['bp']}/10^7", CR, BLACK),
                "tot": (f"={col}{r['svc']}+{col}{r['bm']}", CR, BOLD),
                "ind": (f"={col}{r['tot']}*Inputs!$B$22", CR, BLACK),
                "ka": (f"={col}{r['tot']}*Inputs!$B$23", CR, BLACK),
                "bsh": (f"=IF({col}{r['tot']}=0,0,{col}{r['bm']}/{col}{r['tot']})", PCT, BLACK),
            }
            for key, (f, fmt, font) in cells.items():
                c = ws.cell(r[key], FIRST_COL + k, f); c.number_format = fmt; c.font = font; c.border = THIN
    ws["A47"] = "Green = linked from Inputs; black = calculated. Rs Cr = Rs crore (10 million rupees)."
    ws.column_dimensions["A"].width = 42
    for k in range(len(YEARS)):
        ws.column_dimensions[L(FIRST_COL + k)].width = 12
    return ws


def cell(name, line, year):
    """Address of a model line for a scenario and year, e.g. Model!$L$26."""
    return f"Model!${L(FIRST_COL + YEARS.index(year))}${BLOCK_START[name] + LINES.index(line) + 1}"


def summary(wb):
    ws = wb.create_sheet("Summary")
    ws["A1"], ws["A1"].font = "Summary: India E2W aftermarket", TITLE
    hdr(ws, 3, ["Metric", "Bear", "Base", "Bull"])
    T, F, S, B, I, K = LINES[7], LINES[1], LINES[3], LINES[6], LINES[8], LINES[9]
    metrics = [("Total aftermarket FY26 (Rs Cr)", lambda n: f"={cell(n, T, 2026)}", CR),
               ("Total aftermarket FY31 (Rs Cr)", lambda n: f"={cell(n, T, 2031)}", CR),
               ("CAGR FY26-FY31", lambda n: f"=({cell(n, T, 2031)}/{cell(n, T, 2026)})^(1/5)-1", PCT),
               ("Fleet in use FY31 (million)", lambda n: f"={cell(n, F, 2031)}/10^6", '0.0'),
               ("Battery replacement share FY31", lambda n: f"={cell(n, B, 2031)}/{cell(n, T, 2031)}", PCT),
               ("Independent-garage addressable FY31 (Rs Cr)", lambda n: f"={cell(n, I, 2031)}", CR),
               ("Karnataka market FY31 (Rs Cr)", lambda n: f"={cell(n, K, 2031)}", CR)]
    for i, (label, f, fmt) in enumerate(metrics, start=4):
        ws.cell(i, 1, label)
        for j, n in enumerate(["Bear", "Base", "Bull"], start=2):
            c = ws.cell(i, j, f(n)); c.number_format = fmt
    # chart data: total market by year and scenario
    ws["A13"] = "Total aftermarket by year (Rs Cr)"; ws["A13"].font = BOLD
    hdr(ws, 14, ["Fiscal year", "Bear", "Base", "Bull", "Base: service", "Base: battery"])
    for i, y in enumerate(YEARS, start=15):
        ws.cell(i, 1, f"FY{str(y)[2:]}")
        for j, n in enumerate(["Bear", "Base", "Bull"], start=2):
            ws.cell(i, j, f"={cell(n, T, y)}").number_format = CR
        ws.cell(i, 5, f"={cell('Base', S, y)}").number_format = CR
        ws.cell(i, 6, f"={cell('Base', B, y)}").number_format = CR
    ch = LineChart(); ch.title = "Total E2W aftermarket by scenario (Rs Cr)"; ch.height, ch.width = 8, 16
    ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=14, max_row=24), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=15, max_row=24))
    ws.add_chart(ch, "H3")
    bc = BarChart(); bc.type = "col"; bc.grouping = "stacked"; bc.overlap = 100
    bc.title = "Base case mix: service vs battery (Rs Cr)"; bc.height, bc.width = 8, 16
    bc.add_data(Reference(ws, min_col=5, max_col=6, min_row=14, max_row=24), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=15, max_row=24))
    ws.add_chart(bc, "H20")
    ws.column_dimensions["A"].width = 46
    for col in "BCDEF":
        ws.column_dimensions[col].width = 14
    return ws


def sensitivity(wb):
    ws = wb.create_sheet("Sensitivity")
    ws["A1"], ws["A1"].font = "Sensitivity: FY31 total aftermarket (Rs Cr), other inputs at base", TITLE
    growth = [0.10, 0.15, 0.20, 0.25, 0.30]
    spend = [2000, 2500, 3000, 3500, 4000]
    ws["A3"] = "Helper: fleet by growth rate"; ws["A3"].font = BOLD
    ws["A4"] = "Sales growth"
    for j, g in enumerate(growth, start=2):
        c = ws.cell(4, j, g); c.font, c.number_format = BLUE, PCT
    fy26_sales, fy26_fleet = cell("Base", LINES[0], 2026), cell("Base", LINES[1], 2026)
    for k, y in enumerate(range(2027, 2032)):
        rs, rf = 5 + k, 11 + k
        ws.cell(rs, 1, f"Sales FY{str(y)[2:]}"); ws.cell(rf, 1, f"Fleet FY{str(y)[2:]}")
        for j in range(2, 7):
            col = L(j)
            prev_s = fy26_sales if k == 0 else f"{col}{rs - 1}"
            prev_f = fy26_fleet if k == 0 else f"{col}{rf - 1}"
            ws.cell(rs, j, f"={prev_s}*(1+{col}$4)").number_format = UNITS
            ws.cell(rf, j, f"={prev_f}*(1-Inputs!$B$18)+{col}{rs}").number_format = UNITS
    ws["A16"] = "Paying fleet FY31"
    for j in range(2, 7):
        ws.cell(16, j, f"={L(j)}15-{L(j)}9").number_format = UNITS
    ws["A17"] = "Battery market FY31 (Rs Cr, from FY25 vintage; growth-independent)"
    ws["B17"] = f"={cell('Base', LINES[6], 2031)}"; ws["B17"].number_format = CR
    ws["A19"] = "Spend per vehicle (Rs/yr) ↓  ×  sales growth →"; ws["A19"].font = BOLD
    for j, g in enumerate(growth, start=2):
        c = ws.cell(19, j, f"={L(j)}4"); c.number_format = PCT; c.font = WHITE_B; c.fill = HEAD
    for i, s in enumerate(spend, start=20):
        c = ws.cell(i, 1, s); c.font, c.number_format = BLUE, UNITS
        for j in range(2, 7):
            ws.cell(i, j, f"={L(j)}$16*$A{i}/10^7+$B$17").number_format = CR
    ws.column_dimensions["A"].width = 52
    for j in range(2, 7):
        ws.column_dimensions[L(j)].width = 14


def cover(wb):
    ws = wb.active; ws.title = "Cover"
    ws["A1"], ws["A1"].font = "India electric two-wheeler aftermarket: market sizing, FY22-FY31", TITLE
    lines = [
        ("Question", "How big is the servicing and battery-replacement market for India's electric two-wheelers, and how fast will it grow?"),
        ("Approach", "Bottom-up: vehicles sold each year (Vahan-based sources) → fleet in use after scrappage → paying vehicles x annual service spend, plus battery replacements as each year's sales reach replacement age."),
        ("Cross-check", "Annual spend per vehicle sits inside the Rs 2,000-5,000 reported for e-scooters, about half the Rs 3,500-6,000 for a petrol scooter."),
        ("Base answer", '="Base case: Rs "&TEXT(Summary!C4,"#,##0")&" Cr in FY26, rising to Rs "&TEXT(Summary!C5,"#,##0")&" Cr by FY31 ("&TEXT(Summary!C6,"0%")&" a year)."'),
        ("Range", '="Bear to bull FY31: Rs "&TEXT(Summary!B5,"#,##0")&" Cr to Rs "&TEXT(Summary!D5,"#,##0")&" Cr."'),
        ("Colour code", "Blue = input (sourced); blue on yellow = assumption to challenge; green = link to Inputs; black = formula."),
        ("How to use", "Change blue cells on Inputs; everything else recalculates. Sensitivity shows FY31 size across growth and spend."),
    ]
    for i, (a, b) in enumerate(lines, start=3):
        ws.cell(i, 1, a).font = BOLD
        c = ws.cell(i, 2, b); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A11"] = "Sources"; ws["A11"].font = BOLD
    for i, s in enumerate(SOURCES.values(), start=12):
        ws.cell(i, 2, s).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 120


def build():
    wb = Workbook()
    cover(wb); inputs(wb); model(wb); summary(wb); sensitivity(wb)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.font and c.font.name != "Arial":
                    c.font = Font(name="Arial", bold=c.font.bold, color=c.font.color, size=c.font.size)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    build()
