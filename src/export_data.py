"""Read the calculated workbook and write deck/data.json for the slide deck."""
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
wb = load_workbook(ROOT / "model" / "ev2w_service_market.xlsx", data_only=True)
s, m, se = wb["Summary"], wb["Model"], wb["Sensitivity"]
if s["C5"].value is None:
    raise SystemExit("Workbook has no calculated values: open and save it in Excel/LibreOffice first.")
years = [m.cell(2, c).value for c in range(3, 13)]
row = lambda r: [m.cell(r, c).value for c in range(3, 13)]
data = {
    "years": [f"FY{str(y)[2:]}" for y in years],
    "sales_hist": row(5)[:5],
    "total": {"Bear": row(12), "Base": row(26), "Bull": row(40)},
    "base_service": row(22), "base_battery": row(25), "base_fleet": row(20),
    "summary": {s.cell(r, 1).value: [s.cell(r, c).value for c in (2, 3, 4)] for r in range(4, 11)},
    "sens_growth": [se.cell(4, c).value for c in range(2, 7)],
    "sens_spend": [se.cell(r, 1).value for r in range(20, 25)],
    "sens": [[se.cell(r, c).value for c in range(2, 7)] for r in range(20, 25)],
}
(ROOT / "deck").mkdir(exist_ok=True)
(ROOT / "deck" / "data.json").write_text(json.dumps(data, indent=1))
print("wrote deck/data.json")
